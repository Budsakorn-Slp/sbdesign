"""หน้าพนักงานขาย: รูปตัวโชว์รายสาขา · หมายเหตุ · พนักงานร่วมบิล Z1-ZK · template · export

ด่านที่สำคัญที่สุดคือสิทธิ์ — ทุกเคสยิงผ่าน API จริง ไม่ได้เรียก service ตรงๆ
เพราะสิ่งที่ต้องพิสูจน์คือ "หน้าเว็บส่งอะไรมาก็ข้ามด่านไม่ได้"
"""
import csv
import io
from datetime import date

import pytest
from PIL import Image
from sqlalchemy import select

from app.core.security import hash_password
from app.db.session import SessionLocal
from app.models.photo import ProductPhoto, ProductPhotoAudit
from app.models.quotation import QuotationTemplate
from app.models.user import User
from app.seed import seed_catalog
from app.services import photo_service, sales_extras_service
from app.services.sales_extras_service import month_end
from tests.helpers import auth_headers, ensure_seed

MATNR = "20000158"     # ตัวโชว์ (ขึ้นต้น 20)


@pytest.fixture(autouse=True)
def _setup(tmp_path, monkeypatch):
    ensure_seed()
    with SessionLocal() as db:
        seed_catalog(db)
        # พนักงานอีกสาขา — ใช้พิสูจน์ว่าเห็นรูปข้ามสาขาไม่ได้
        if not db.scalar(select(User).where(User.staff_code == "SA-900")):
            db.add(User(role="sales", name="สมปอง อีกสาขา", staff_code="SA-900", branch_id="S319",
                        email="sompong@sb.local", password_hash=hash_password("1122"), is_guest=False))
            db.commit()
        # เทสไฟล์อื่น (cli.secure set-branch) ย้าย/ล้างสาขาของบัญชีร่วมได้ — ตั้งกลับให้แน่นอน
        for u in db.scalars(select(User).where(User.staff_code.in_(("SA-104", "SA-105", "MG-001")))):
            u.branch_id = "BKN"
        db.commit()
    # รูปที่อัปโหลดในเทสไปลงโฟลเดอร์ชั่วคราว ไม่ปนกับรูปจริงของเครื่อง dev
    monkeypatch.setattr(photo_service, "PHOTO_DIR", tmp_path / "product_photos")
    monkeypatch.setattr(sales_extras_service, "UPLOAD_ROOT", tmp_path)
    yield
    # template ผูกกับบัญชี SA-104 ที่เทสไฟล์อื่นใช้ร่วม — ทิ้งไว้จะไปทับเบอร์/เงื่อนไขในใบของเทสอื่น
    with SessionLocal() as db:
        db.query(QuotationTemplate).delete()
        db.commit()


def _jpeg(color="red", exif_gps=False) -> bytes:
    im = Image.new("RGB", (64, 48), color)
    buf = io.BytesIO()
    if exif_gps:
        ex = Image.Exif()
        ex[0x8825] = {1: "N", 2: (13.0, 45.0, 0.0)}   # GPSInfo
        ex[0x0110] = "Pixel 9"                       # Model
        im.save(buf, format="JPEG", exif=ex)
    else:
        im.save(buf, format="JPEG")
    return buf.getvalue()


def _upload(client, hs, n=1, matnr=MATNR):
    files = [("files", (f"p{i}.jpg", _jpeg(), "image/jpeg")) for i in range(n)]
    return client.post(f"/staff/materials/{matnr}/photos", files=files, headers=hs)


# ---------- รูปตัวโชว์ ----------
def test_อัปโหลดหลายรูป_สาขามาจากบัญชีไม่ใช่จากคำขอ(client):
    hs = auth_headers(client, "SA-104", "staff")
    r = client.post(f"/staff/materials/{MATNR}/photos?branch_code=S319",
                    files=[("files", ("a.jpg", _jpeg(), "image/jpeg")), ("files", ("b.jpg", _jpeg("blue"), "image/jpeg"))],
                    data={"branch_code": "S319"}, headers=hs)
    assert r.status_code == 201, r.text
    out = r.json()
    assert len(out) == 2
    assert all(p["branch_code"] == "BKN" for p in out), "ต้องใช้สาขาของบัญชี ไม่ใช่ที่หน้าเว็บส่งมา"
    assert all(p["mine"] and p["can_edit"] and p["can_delete"] for p in out)
    assert out[0]["owner_employee_code"] == "SA-104"


def test_ถ่ายรูปได้เฉพาะตัวโชว์(client):
    hs = auth_headers(client, "SA-104", "staff")
    assert _upload(client, hs, matnr="10023841").status_code == 400


def test_ไฟล์ปลอมนามสกุล_jpg_ไม่ผ่าน(client):
    hs = auth_headers(client, "SA-104", "staff")
    r = client.post(f"/staff/materials/{MATNR}/photos",
                    files=[("files", ("evil.jpg", b"<script>alert(1)</script>", "image/jpeg"))], headers=hs)
    assert r.status_code == 400


def test_ล้าง_GPS_ออกจากรูป(client):
    hs = auth_headers(client, "SA-104", "staff")
    r = client.post(f"/staff/materials/{MATNR}/photos",
                    files=[("files", ("gps.jpg", _jpeg(exif_gps=True), "image/jpeg"))], headers=hs)
    assert r.status_code == 201, r.text
    with SessionLocal() as db:
        path = db.get(ProductPhoto, r.json()[0]["id"]).file_path
    saved = Image.open(photo_service.PHOTO_DIR / path.split("/")[-1])
    ex = saved.getexif()
    assert 0x8825 not in ex and 0x0110 not in ex


def test_เพื่อนร่วมสาขาเห็นแต่แก้ลบไม่ได้(client):
    a = auth_headers(client, "SA-104", "staff")
    b = auth_headers(client, "SA-105", "staff")    # สาขาเดียวกัน
    pid = _upload(client, a).json()[0]["id"]
    seen = {p["id"]: p for p in client.get(f"/staff/materials/{MATNR}/photos", headers=b).json()}
    assert pid in seen and not seen[pid]["can_edit"] and not seen[pid]["can_delete"]
    assert client.delete(f"/staff/photos/{pid}", headers=b).status_code == 403
    r = client.put(f"/staff/photos/{pid}", files={"file": ("x.jpg", _jpeg(), "image/jpeg")}, headers=b)
    assert r.status_code == 403


def test_ต่างสาขามองไม่เห็น_และตอบ404ไม่ใช่403(client):
    a = auth_headers(client, "SA-104", "staff")
    other = auth_headers(client, "SA-900", "staff")
    pid = _upload(client, a).json()[0]["id"]
    ids = [p["id"] for p in client.get(f"/staff/materials/{MATNR}/photos?branch=BKN", headers=other).json()]
    assert pid not in ids, "ขอสาขาอื่นผ่าน query ต้องไม่ได้ผล"
    # 404 = ไม่บอกใบ้ว่ารูปของสาขาอื่นมีอยู่จริง
    assert client.delete(f"/staff/photos/{pid}", headers=other).status_code == 404


def test_เจ้าของแทนรูปและลบได้_ลบแบบซ่อน_ประวัติครบ(client):
    a = auth_headers(client, "SA-104", "staff")
    pid = _upload(client, a).json()[0]["id"]
    r = client.put(f"/staff/photos/{pid}", files={"file": ("n.jpg", _jpeg("green"), "image/jpeg")}, headers=a)
    assert r.status_code == 200, r.text
    assert client.delete(f"/staff/photos/{pid}", headers=a).status_code == 204
    assert pid not in [p["id"] for p in client.get(f"/staff/materials/{MATNR}/photos", headers=a).json()]
    with SessionLocal() as db:
        row = db.get(ProductPhoto, pid)
        assert row.is_deleted and row.deleted_by_employee == "SA-104", "ต้องเป็น soft delete"
        acts = [x.action for x in db.scalars(select(ProductPhotoAudit).where(ProductPhotoAudit.image_id == pid)
                                              .order_by(ProductPhotoAudit.action_at)).all()]
    assert acts == ["CREATE", "UPDATE", "DELETE"]


def test_ผู้จัดการลบรูปคนอื่นได้_ประวัติแยกเจ้าของกับคนทำ(client):
    a = auth_headers(client, "SA-104", "staff")
    m = auth_headers(client, "MG-001", "staff")
    pid = _upload(client, a).json()[0]["id"]
    assert client.delete(f"/staff/photos/{pid}", headers=m).status_code == 204
    log = [x for x in client.get(f"/staff/photos/audit?matnr={MATNR}", headers=m).json()
           if x["image_id"] == pid and x["action"] == "DELETE"]
    assert log and log[0]["image_owner_employee"] == "SA-104"
    assert log[0]["action_by_employee"] == "MG-001" and log[0]["action_by_role"] == "manager"


def test_พนักงานขายดูประวัติไม่ได้(client):
    a = auth_headers(client, "SA-104", "staff")
    assert client.get("/staff/photos/audit", headers=a).status_code == 403


def test_ลูกค้าเข้าไม่ได้(client):
    c = auth_headers(client, "1100440182", "customer")
    assert client.get(f"/staff/materials/{MATNR}/photos", headers=c).status_code == 403


# ---------- ตะกร้า: หมายเหตุ + พนักงานร่วมบิล ----------
def _new_cart(client, hs):
    cart = client.post("/sales/carts", json={}, headers=hs).json()
    client.post(f"/sales/carts/{cart['id']}/items", json={"matnr": "10023841", "qty": 1, "supply_mode": "ship"}, headers=hs)
    return cart


def _emp(client, hs, code):
    return next(e for e in client.get("/staff/employees", headers=hs).json() if e["employee_code"] == code)


def test_dropdown_พนักงานเป็นรหัส_ชื่อ(client):
    hs = auth_headers(client, "SA-104", "staff")
    e = _emp(client, hs, "SA-105")
    assert e["label"] == "SA-105 – สมหญิง ข."
    assert [r["code"] for r in client.get("/staff/roles", headers=hs).json()] == ["Z1", "Z2", "Z3", "Z4", "ZK"]


def test_ผูกพนักงานรายบทบาท_ชื่อมาจากทะเบียน(client):
    hs = auth_headers(client, "SA-104", "staff")
    cart = _new_cart(client, hs)
    e = _emp(client, hs, "SA-105")
    r = client.put(f"/sales/carts/{cart['id']}/staff",
                   json={"role_code": "Z1", "user_id": e["user_id"], "employee_name": "ชื่อปลอม"}, headers=hs)
    assert r.status_code == 200, r.text
    staff = r.json()["staff"]
    assert staff == [{"role_code": "Z1", "role_name": staff[0]["role_name"], "user_id": e["user_id"],
                      "employee_code": "SA-105", "employee_name": "สมหญิง ข."}]
    # เปลี่ยนคนในบทบาทเดิม = แทนที่ ไม่ใช่เพิ่มแถว
    me = _emp(client, hs, "SA-104")
    staff = client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z1", "user_id": me["user_id"]}, headers=hs).json()["staff"]
    assert len(staff) == 1 and staff[0]["employee_code"] == "SA-104"
    # ถอดออก
    staff = client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z1", "user_id": None}, headers=hs).json()["staff"]
    assert staff == []


def test_บทบาทที่ไม่รู้จักและพนักงานที่ไม่มีจริงไม่ผ่าน(client):
    hs = auth_headers(client, "SA-104", "staff")
    cart = _new_cart(client, hs)
    assert client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z9", "user_id": None}, headers=hs).status_code == 422
    assert client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z3", "user_id": "ghost"}, headers=hs).status_code == 404


def test_แก้ตะกร้าของเซลล์คนอื่นไม่ได้(client):
    a = auth_headers(client, "SA-104", "staff")
    other = auth_headers(client, "SA-900", "staff")
    cart = _new_cart(client, a)
    assert client.put(f"/sales/carts/{cart['id']}/remark", json={"overall_remark": "x"}, headers=other).status_code == 403


def test_หมายเหตุหลักกับหมายเหตุรายสินค้าแยกกัน_และติดไปกับใบ(client):
    from tests.test_step8_quotation import _ready_cart

    hs = auth_headers(client, "SA-104", "staff")
    cart, _ = _ready_cart(client, hs)
    item = cart["items"][0]
    client.patch(f"/sales/carts/{cart['id']}/items/{item['id']}", json={"note": "เปลี่ยนผ้าเป็นสีเทา"}, headers=hs)
    r = client.put(f"/sales/carts/{cart['id']}/remark", json={"overall_remark": "ส่งก่อน 10 โมง"}, headers=hs)
    assert r.status_code == 200 and r.json()["overall_remark"] == "ส่งก่อน 10 โมง"
    z = _emp(client, hs, "SA-105")
    client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z2", "user_id": z["user_id"]}, headers=hs)
    client.put("/staff/quotation-template", json={"display_name": "สมชาย (DS)", "bank_accounts": "กสิกร 111-1-11111-1",
                                                  "footer_terms": "ราคานี้ถึง {month_end}", "standard_remark": "ราคารวม VAT"}, headers=hs)

    p = client.post("/presos", json={"cart_id": cart["id"]}, headers=hs).json()
    r = client.post(f"/presos/{p['preso_no']}/quotation", json={"force": True}, headers=hs)
    assert r.status_code == 201, (p, r.text)
    q = r.json()
    doc = client.get(f"/quotations/{q['quotation_no']}/document", headers=hs).text
    assert "ส่งก่อน 10 โมง" in doc and "เปลี่ยนผ้าเป็นสีเทา" in doc and "ราคารวม VAT" in doc
    assert "SA-105 – สมหญิง ข." in doc and "สมชาย (DS)" in doc and "กสิกร 111-1-11111-1" in doc
    issued = date.fromisoformat(q["issued_at"][:10])
    assert f"ราคานี้ถึง {month_end(issued).strftime('%d/%m/%Y')}" in doc

    # แก้ template ทีหลัง ใบที่ออกไปแล้วต้องไม่เปลี่ยน
    client.put("/staff/quotation-template", json={"display_name": "ชื่อใหม่"}, headers=hs)
    doc2 = client.get(f"/quotations/{q['quotation_no']}/document", headers=hs).text
    assert "สมชาย (DS)" in doc2 and "ชื่อใหม่" not in doc2

    # export ทั้งสองแบบมีข้อมูลชุดเดียวกัน
    r = client.get(f"/quotations/{q['quotation_no']}/export?format=csv", headers=hs)
    assert r.status_code == 200 and r.content.startswith("﻿".encode())
    text = r.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    flat = " ".join(" ".join(x) for x in rows)
    for must in ("เปลี่ยนผ้าเป็นสีเทา", "ส่งก่อน 10 โมง", "กสิกร 111-1-11111-1", "SA-105 – สมหญิง ข.",
                 month_end(issued).strftime("%d/%m/%Y"), item["matnr"]):
        assert must in flat, must
    r = client.get(f"/quotations/{q['quotation_no']}/export?format=xlsx", headers=hs)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).active
    cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "เปลี่ยนผ้าเป็นสีเทา" in cells and "สมชาย (DS)" in cells
    # คนที่เปิดใบไม่ได้ ก็ export ไม่ได้
    assert client.get(f"/quotations/{q['quotation_no']}/export?format=csv",
                      headers=auth_headers(client, "SA-900", "staff")).status_code == 403


# ---------- template ----------
def test_template_เป็นของตัวเองเท่านั้น(client):
    a = auth_headers(client, "SA-104", "staff")
    b = auth_headers(client, "SA-105", "staff")
    client.put("/staff/quotation-template", json={"phone": "081-111-1111"}, headers=a)
    assert client.get("/staff/quotation-template", headers=b).json()["phone"] != "081-111-1111"
    assert client.get("/staff/quotation-template", headers=a).json()["phone"] == "081-111-1111"


def test_อัปโหลดโลโก้(client):
    a = auth_headers(client, "SA-104", "staff")
    r = client.post("/staff/quotation-template/logo", files={"file": ("logo.png", _jpeg(), "image/png")}, headers=a)
    assert r.status_code == 200, r.text
    assert r.json()["logo_url"].startswith("/media/product_photos/")


@pytest.mark.parametrize("d,want", [(date(2026, 10, 6), date(2026, 10, 31)), (date(2026, 11, 15), date(2026, 11, 30)),
                                    (date(2028, 2, 3), date(2028, 2, 29)), (date(2026, 12, 31), date(2026, 12, 31))])
def test_วันสุดท้ายของเดือน(d, want):
    assert month_end(d) == want


# ---------- แชร์รูปให้ลูกค้าเห็น ----------
def test_อัปโหลดแล้วลูกค้ายังไม่เห็นจนกว่าจะกดแชร์(client):
    a = auth_headers(client, "SA-104", "staff")
    pid = _upload(client, a).json()[0]["id"]
    seen = lambda: [p["id"] for b in client.get(f"/materials/{MATNR}/branch-photos").json() for p in b["photos"]]
    assert pid not in seen(), "รูปที่ยังไม่แชร์ต้องไม่หลุดถึงลูกค้า"
    r = client.post(f"/staff/photos/{pid}/share", json={"public": True}, headers=a)
    assert r.status_code == 200 and r.json()["is_public"]
    assert pid in seen()
    group = next(b for b in client.get(f"/materials/{MATNR}/branch-photos").json() if any(p["id"] == pid for p in b["photos"]))
    assert group["branch_code"] == "BKN"
    assert "owner_employee_code" not in str(group), "ไม่ส่งข้อมูลพนักงานให้ลูกค้า"
    client.post(f"/staff/photos/{pid}/share", json={"public": False}, headers=a)
    assert pid not in seen()
    # ลบรูปที่แชร์อยู่ = ลูกค้าไม่เห็นทันที
    client.post(f"/staff/photos/{pid}/share", json={"public": True}, headers=a)
    client.delete(f"/staff/photos/{pid}", headers=a)
    assert pid not in seen()
    with SessionLocal() as db:
        acts = [x.action for x in db.scalars(select(ProductPhotoAudit).where(ProductPhotoAudit.image_id == pid)
                                              .order_by(ProductPhotoAudit.action_at)).all()]
    assert acts == ["CREATE", "SHARE", "UNSHARE", "SHARE", "DELETE"]


def test_แชร์รูปของเพื่อนไม่ได้(client):
    a = auth_headers(client, "SA-104", "staff")
    b = auth_headers(client, "SA-105", "staff")
    pid = _upload(client, a).json()[0]["id"]
    assert client.post(f"/staff/photos/{pid}/share", json={"public": True}, headers=b).status_code == 403


def test_ใส่พนักงานร่วมบิลด้วยรหัส(client):
    hs = auth_headers(client, "SA-104", "staff")
    cart = _new_cart(client, hs)
    r = client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z3", "employee_code": " sa-105 "}, headers=hs)
    assert r.status_code == 200, r.text
    assert r.json()["staff"][0]["employee_name"] == "สมหญิง ข."
    r = client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z4", "employee_code": "XX-999"}, headers=hs)
    assert r.status_code == 404 and "XX-999" in r.text
    r = client.put(f"/sales/carts/{cart['id']}/staff", json={"role_code": "Z3", "employee_code": "ADM-001"}, headers=hs)
    assert r.status_code == 404, "แอดมินระบบไม่ใช่คนขาย"


# ---------- Google Sheets ----------
def _issued(client, hs):
    from tests.test_step8_quotation import _ready_cart

    cart, _ = _ready_cart(client, hs)
    p = client.post("/presos", json={"cart_id": cart["id"]}, headers=hs).json()
    r = client.post(f"/presos/{p['preso_no']}/quotation", json={"force": True}, headers=hs)
    assert r.status_code == 201, r.text
    return r.json()


def test_google_sheet_ยังไม่ตั้งค่าตอบ501_และ_csv_ไม่มี_bom(client):
    hs = auth_headers(client, "SA-104", "staff")
    q = _issued(client, hs)
    assert client.post(f"/quotations/{q['quotation_no']}/google-sheet", headers=hs).status_code == 501
    r = client.get(f"/quotations/{q['quotation_no']}/export?format=csv&bom=0", headers=hs)
    assert r.status_code == 200 and not r.content.startswith("﻿".encode())


def test_google_sheet_สร้างชีตและแชร์ให้คนกด(client, tmp_path, monkeypatch):
    import json

    import httpx
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa

    from app.core.config import get_settings
    from app.integrations.google import sheets

    pem = rsa.generate_private_key(public_exponent=65537, key_size=2048).private_bytes(
        serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode()
    key = tmp_path / "sa.json"
    key.write_text(json.dumps({"client_email": "bot@x.iam.gserviceaccount.com", "private_key": pem}))
    s = get_settings()
    monkeypatch.setattr(s, "google_sa_file", str(key))
    monkeypatch.setattr(s, "google_sheets_folder_id", "FOLDER1")
    calls = []

    def handler(req: httpx.Request) -> httpx.Response:
        calls.append((req.method, req.url.path, req.content))
        if req.url.path == "/token":
            return httpx.Response(200, json={"access_token": "tok"})
        if req.url.path == "/drive/v3/files":
            assert json.loads(req.content)["parents"] == ["FOLDER1"]
            return httpx.Response(200, json={"id": "SHEET1"})
        return httpx.Response(200, json={})

    monkeypatch.setattr(sheets, "_client", lambda: httpx.Client(transport=httpx.MockTransport(handler)))
    hs = auth_headers(client, "SA-104", "staff")
    q = _issued(client, hs)
    r = client.post(f"/quotations/{q['quotation_no']}/google-sheet", headers=hs)
    assert r.status_code == 200, r.text
    assert r.json()["url"] == "https://docs.google.com/spreadsheets/d/SHEET1/edit"
    put = next(c for c in calls if c[0] == "PUT")
    assert q["quotation_no"] in put[2].decode()
    share = next(c for c in calls if c[1].endswith("/permissions"))
    assert json.loads(share[2])["emailAddress"] == "somchai@sb.local"
    # เซลล์คนอื่นที่ไม่ใช่เจ้าของใบ สร้างชีตของใบนี้ไม่ได้
    other = auth_headers(client, "SA-900", "staff")
    assert client.post(f"/quotations/{q['quotation_no']}/google-sheet", headers=other).status_code == 403
