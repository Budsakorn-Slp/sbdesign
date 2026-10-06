"""ใบเสนอราคาเป็นไฟล์ตาราง — Excel (.xlsx) · CSV · Google Sheets

เนื้อหาชุดเดียวกับใบ PDF (api/quotation_doc.py) ทุกอย่าง:
  หัวใบ (โลโก้ บริษัท เลขที่/วันที่/สาขา) · ลูกค้ากับสถานที่ส่งคู่กัน · รหัสลูกค้า/พนักงานขาย
  · หมายเหตุ · ตารางสินค้า · ยอดรวม · บัญชีโอน · เงื่อนไขท้ายใบ

ไม่มีแถวพนักงานร่วมบิล Z1-ZK — ไฟล์นี้ส่งต่อให้ลูกค้าได้ ข้อมูลค่าคอมไม่ควรติดไป
บรรทัดค่าบริการ (เปิด Mat ค่าขนส่ง เช่น A534) ไม่อยู่ในตารางสินค้า — ยอดอยู่ที่ช่องค่าขนส่งแล้ว
ไม่งั้นลูกค้าเห็นค่าส่งสองที่ (แบบเดียวกับที่ใบ PDF แก้ไปแล้ว)
"""
from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from app.api.quotation_doc import COMPANY_LOGO, STATUS_TXT, SUPPLY_TXT, _terms, month_end_of
from app.core.config import get_settings
from app.models.quotation import Quotation
from app.services import sales_extras_service
from app.services.photo_service import UPLOAD_ROOT

log = logging.getLogger("sb.export")

HEAD = ["ลำดับ", "รหัสสินค้า", "รายการ", "จำนวน", "ราคาต่อหน่วย", "ส่วนลด", "จำนวนเงิน", "รับสินค้า", "หมายเหตุรายสินค้า"]


@dataclass
class Sheet:
    company: list[str]
    meta: list[tuple[str, str]]
    bill: list[str]
    ship: list[str]
    who: str
    remark: str
    items: list[list]
    totals: list[tuple[str, float]]
    bank: str
    terms: list[str]
    logo_path: str | None = None


def build(q: Quotation) -> Sheet:
    from app.services import staff_shipping_service

    s = get_settings()
    tp = q.template_snapshot or {}
    c = q.customer_snapshot or {}
    charges = staff_shipping_service.charge_matnrs()

    b_code = (tp.get("branch_code") or (q.sales.branch_id if q.sales else None) or "").lstrip("S")
    b_name = tp.get("branch_name") or (q.sales.branch_name if q.sales else None)
    meta = []
    if b_name:
        meta.append(("สาขา", f"{b_code} · {b_name}".strip(" ·")))
    meta += [("เลขที่ใบเสนอราคา", q.quotation_no), ("วันที่ออก", q.issued_at.strftime("%d/%m/%Y")),
             ("ยืนราคาถึง", q.valid_until.strftime("%d/%m/%Y")), ("สถานะ", STATUS_TXT.get(q.status, q.status))]

    bill_addr = " ".join(x for x in [c.get("address") or "", c.get("postcode") or ""] if x).strip()
    ship_addr = " ".join(x for x in [q.ship_address or "", q.ship_postcode or ""] if x).strip()
    phone_mail = (c.get("phone") or "") + ((" · " + c["email"]) if c.get("email") else "")
    bill = [c.get("name") or "", bill_addr or ship_addr or "-", phone_mail]
    ship = [c.get("name") or "", ship_addr or "ยังไม่ได้ระบุ", f"Tel. {c['phone']}" if c.get("phone") else ""]

    sales_name = tp.get("employee_name") or (q.sales.name if q.sales else "สั่งซื้อออนไลน์")
    code = q.sales.staff_code if q.sales and q.sales.staff_code else ""
    phone = tp.get("phone") or (q.sales.phone if q.sales else None)
    who = f"รหัสลูกค้า  {c.get('sap_customer_no') or '-'}          พนักงานขาย  {sales_name}{f' ({code})' if code else ''}{f' · โทร {phone}' if phone else ''}"

    items = []
    n = 0
    for l in q.lines:
        if l.matnr in charges:
            continue
        n += 1
        off = (Decimal(l.list_price) - Decimal(l.unit_price)) * l.qty
        items.append([n, l.matnr, l.name + (f" ({l.variant})" if l.variant else ""), l.qty, float(l.unit_price),
                      float(off) if off > 0 else 0.0, float(l.line_total), SUPPLY_TXT.get(l.supply_mode, l.supply_mode or ""),
                      l.item_remark or ""])

    fee = float(q.shipping_fee) + float(q.install_fee) - float(q.shipping_discount)
    totals = [("รวมสินค้า", float(q.subtotal))]
    if float(q.discount_total) > 0:
        totals.append(("ส่วนลด", -float(q.discount_total)))
    if fee:
        totals.append(("ค่าขนส่ง" + (" + ติดตั้ง" if float(q.install_fee) else ""), fee))
    totals.append((f"รวมสุทธิ (รวม VAT 7% = {float(q.vat):,.2f})", float(q.grand_total)))

    remark = "\n".join(x for x in [(q.overall_remark or (q.preso.note if q.preso else None) or "").strip(),
                                   (tp.get("standard_remark") or "").strip()] if x)
    return Sheet(
        company=[s.company_name_th, s.company_name_en, s.company_address_th, s.company_address_en,
                 f"เลขประจำตัวผู้เสียภาษี TAX ID {s.company_tax_id}"],
        meta=meta, bill=bill, ship=ship, who=who, remark=remark, items=items, totals=totals,
        bank=(tp.get("bank_accounts") or "").strip(),
        terms=_terms(q.valid_until.strftime("%d/%m/%Y"), month_end_of(q), tp.get("footer_terms")),
        logo_path=tp.get("logo_path"),
    )


# ---------- แถวเรียบ (CSV / Google Sheets) ----------
def flat_rows(q: Quotation) -> list[list]:
    d = build(q)
    rows: list[list] = [["ใบเสนอราคา / Quotation"], *[[x] for x in d.company], []]
    rows += [[k, v] for k, v in d.meta]
    rows += [[], ["ชื่อ-ที่อยู่ลูกค้า", "", "", "", "ชื่อ-สถานที่ส่งสินค้า"]]
    rows += [[a, "", "", "", b] for a, b in zip(d.bill, d.ship)]
    rows += [[], [d.who]]
    if d.remark:
        rows += [[], ["หมายเหตุ", d.remark]]
    rows += [[], HEAD, *d.items, []]
    rows += [["", "", "", "", "", k, round(v, 2)] for k, v in d.totals]
    if d.bank:
        rows += [[], ["บัญชีสำหรับโอนชำระ", d.bank]]
    rows += [[], ["เงื่อนไข"], *[[i, t] for i, t in enumerate(d.terms, 1)]]
    return rows


def csv_bytes(q: Quotation, bom: bool = True) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in flat_rows(q):
        w.writerow([f"{v:.2f}" if isinstance(v, float) else v for v in r])
    # BOM ให้ Excel ภาษาไทยเปิดแล้วไม่เป็นตัวต่างดาว · IMPORTDATA ของ Google ไม่ต้องใช้ (จะติดมาเป็นอักษรล่องหน)
    return (("﻿" if bom else "") + buf.getvalue()).encode("utf-8")


# ---------- Excel ----------
_COMPANY_LOGO_FILE = UPLOAD_ROOT / "_company_logo.png"
_logo_failed = False


def _company_logo() -> Path | None:
    """โลโก้บริษัทเก็บไว้ในเครื่องครั้งแรกที่ใช้ — ไม่ต้องดึงจากเว็บทุกครั้งที่ export
    ดึงไม่ได้ (เครื่องออกเน็ตไม่ได้) = ไม่ใส่โลโก้ ไม่ทำให้ export ล้ม"""
    global _logo_failed
    if _COMPANY_LOGO_FILE.is_file():
        return _COMPANY_LOGO_FILE
    if _logo_failed:
        return None
    try:
        import httpx

        r = httpx.get(COMPANY_LOGO, timeout=4.0, follow_redirects=False)
        if r.status_code == 200 and r.headers.get("content-type", "").startswith("image/"):
            _COMPANY_LOGO_FILE.parent.mkdir(parents=True, exist_ok=True)
            _COMPANY_LOGO_FILE.write_bytes(r.content)
            return _COMPANY_LOGO_FILE
    except Exception as e:  # noqa: BLE001
        log.info("ดึงโลโก้บริษัทไม่สำเร็จ ข้ามโลโก้: %s", e)
    _logo_failed = True
    return None


def _lines_needed(text: str, chars_per_line: int) -> int:
    return sum(max(1, -(-len(part) // chars_per_line)) for part in (text or "").split("\n")) or 1


def xlsx_bytes(q: Quotation) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    d = build(q)
    wb = Workbook()
    ws = wb.active
    ws.title = q.quotation_no[:31]
    widths = {"A": 7, "B": 13, "C": 44, "D": 8, "E": 14, "F": 12, "G": 15, "H": 11, "I": 24}
    for col, wdt in widths.items():
        ws.column_dimensions[col].width = wdt
    base = "Tahoma"   # มีภาษาไทยทุกเครื่อง Windows/Mac — ฟอนต์ตั้งต้นของ Excel บางเครื่องแสดงไทยเพี้ยน
    F = lambda **k: Font(name=base, size=k.pop("size", 10), **k)  # noqa: E731
    gray = "777777"
    thin, thick = Side(style="thin", color="D6D6D4"), Side(style="medium", color="111111")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    def put(cell, value, font=None, align=None, merge=None, fmt=None, fill=None, border=None):
        if merge:
            ws.merge_cells(merge)
        c = ws[cell]
        c.value = value
        c.font = font or F()
        if align:
            c.alignment = align
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if border:
            c.border = border
        return c

    # ---- หัวใบ: โลโก้ + บริษัท (ซ้าย) · ชื่อใบ + เลขที่/วันที่ (ขวา) ----
    logo = sales_extras_service.logo_file(d.logo_path) or _company_logo()
    if logo:
        try:
            from openpyxl.drawing.image import Image as XLImage

            img = XLImage(str(logo))
            ratio = min(90 / max(img.width, 1), 54 / max(img.height, 1))
            img.width, img.height = int(img.width * ratio), int(img.height * ratio)
            ws.add_image(img, "A1")
        except Exception as e:  # noqa: BLE001
            log.info("ใส่โลโก้ใน Excel ไม่ได้: %s", e)
    for i, line in enumerate(d.company, 1):
        put(f"C{i}", line, F(bold=i <= 2, size=10 if i <= 2 else 9, color=None if i <= 2 else "333333"))
    put("G1", "ใบเสนอราคา", F(bold=True, size=16), Alignment(horizontal="right"), merge="G1:I1")
    put("G2", "Quotation", F(color=gray, size=9), Alignment(horizontal="right"), merge="G2:I2")
    for i, (k, v) in enumerate(d.meta, 3):
        put(f"G{i}", k, F(color=gray, size=9), Alignment(horizontal="right"))
        put(f"H{i}", v, F(bold=k in ("เลขที่ใบเสนอราคา", "ยืนราคาถึง"), size=9), Alignment(horizontal="right"), merge=f"H{i}:I{i}")
    row = max(len(d.company), len(d.meta) + 2) + 1
    for col in "ABCDEFGHI":
        ws[f"{col}{row}"].border = Border(bottom=thick)
    row += 2

    # ---- ลูกค้า · สถานที่ส่ง คู่กัน ----
    put(f"A{row}", "ชื่อ-ที่อยู่ลูกค้า", F(bold=True), merge=f"A{row}:D{row}")
    put(f"E{row}", "ชื่อ-สถานที่ส่งสินค้า", F(bold=True), merge=f"E{row}:I{row}")
    for a, b in zip(d.bill, d.ship):
        row += 1
        put(f"A{row}", a, F(), wrap_top, merge=f"A{row}:D{row}")
        put(f"E{row}", b, F(), wrap_top, merge=f"E{row}:I{row}")
        ws.row_dimensions[row].height = 14 * max(_lines_needed(a, 60), _lines_needed(b, 80))
    row += 2
    put(f"A{row}", d.who, F(), Alignment(horizontal="center"), merge=f"A{row}:I{row}")
    for col in "ABCDEFGHI":
        ws[f"{col}{row}"].border = Border(top=thin, bottom=thin)
    row += 1

    # ---- หมายเหตุ (มีเมื่อใส่เท่านั้น) ----
    if d.remark:
        row += 1
        put(f"A{row}", "หมายเหตุ", F(bold=True), wrap_top)
        put(f"B{row}", d.remark, F(), wrap_top, merge=f"B{row}:I{row}")
        ws.row_dimensions[row].height = 14 * _lines_needed(d.remark, 130)

    # ---- ตารางสินค้า ----
    row += 2
    hdr_fill = PatternFill("solid", fgColor="F2F2F0")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    for j, h in enumerate(HEAD):
        col = "ABCDEFGHI"[j]
        put(f"{col}{row}", h, F(bold=True, size=9), Alignment(horizontal="right" if j in (3, 4, 5, 6) else "left", vertical="center", wrap_text=True),
            fill=hdr_fill, border=box)
    for it in d.items:
        row += 1
        for j, v in enumerate(it):
            col = "ABCDEFGHI"[j]
            right = j in (3, 4, 5, 6)
            put(f"{col}{row}", v if not (j == 5 and v == 0) else "-", F(), Alignment(horizontal="right" if right else "left", vertical="top", wrap_text=j in (2, 8)),
                fmt="#,##0.00" if j in (4, 5, 6) and v != 0 else None, border=box)
        ws.row_dimensions[row].height = 14 * max(_lines_needed(str(it[2]), 48), _lines_needed(str(it[8]), 26))

    # ---- ยอดรวม ชิดขวา ----
    row += 1
    for k, v in d.totals:
        row += 1
        last = k.startswith("รวมสุทธิ")
        put(f"C{row}", k, F(bold=last), Alignment(horizontal="right"), merge=f"C{row}:F{row}")
        c = put(f"G{row}", v, F(bold=last), Alignment(horizontal="right"), fmt="#,##0.00")
        if last:
            for col in "CDEFG":
                ws[f"{col}{row}"].border = Border(top=thick)
            c.border = Border(top=thick)

    # ---- บัญชีโอน · เงื่อนไข ----
    if d.bank:
        row += 2
        put(f"A{row}", "บัญชีสำหรับโอนชำระ", F(bold=True), merge=f"A{row}:I{row}")
        row += 1
        put(f"A{row}", d.bank, F(size=9), wrap_top, merge=f"A{row}:I{row}")
        ws.row_dimensions[row].height = 13 * _lines_needed(d.bank, 140)
    row += 2
    put(f"A{row}", "เงื่อนไข", F(bold=True, size=9))
    for i, t in enumerate(d.terms, 1):
        row += 1
        put(f"A{row}", f"{i}.", F(size=9), Alignment(horizontal="right", vertical="top"))
        put(f"B{row}", t, F(size=9), wrap_top, merge=f"B{row}:I{row}")
        ws.row_dimensions[row].height = 13 * _lines_needed(t, 140)

    # พิมพ์พอดีหน้า A4 แนวตั้ง กว้างเต็มหน้า
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
