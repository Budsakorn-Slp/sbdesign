"""เครื่องมือพนักงานขาย: พนักงานร่วมบิล Z1-ZK · หมายเหตุหลักของตะกร้า · template ใบเสนอราคา · export

สิทธิ์ตรวจใน sales_extras_service (ตาม permission) และ sales_service.require_my_cart (ตะกร้าของใคร)
ตรงนี้แค่แปลงคำขอ/คำตอบ
"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from decimal import Decimal

import httpx
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.api.cart import cart_out
from app.api.deps import get_current_user_optional
from app.api.quotation_doc import SUPPLY_TXT, _terms, month_end_of
from app.db.session import get_db
from app.models.cart import STAFF_ROLES
from app.models.quotation import QuotationTemplate
from app.models.user import User
from app.schemas.cart import CartOut
from app.services import audit_service, quotation_service, sales_extras_service, sales_service
from app.services.employee_provider import CurrentEmployee, get_current_employee, get_employee_provider

router = APIRouter(tags=["staff-sales"])


# ---------- schemas ----------
class EmployeeOut(BaseModel):
    user_id: str
    employee_code: str
    employee_name: str
    branch_code: str | None = None
    label: str          # "รหัส – ชื่อ" ใช้เป็นข้อความใน dropdown ตรงๆ


class RoleOut(BaseModel):
    code: str
    name: str


class StaffSetIn(BaseModel):
    role_code: str = Field(min_length=2, max_length=4)
    user_id: str | None = None      # เลือกจากรายชื่อ
    employee_code: str | None = Field(default=None, max_length=32)   # หรือพิมพ์รหัสพนักงาน
    # ทั้งคู่ว่าง = ถอดคนออกจากบทบาทนี้ · ชื่อไม่รับจากหน้าเว็บ — ระบบเปิดทะเบียนหาเองเสมอ


class RemarkIn(BaseModel):
    overall_remark: str | None = Field(default=None, max_length=sales_extras_service.REMARK_MAX)


class TemplateIn(BaseModel):
    display_name: str | None = Field(default=None, max_length=120)
    phone: str | None = Field(default=None, max_length=32)
    bank_accounts: str | None = Field(default=None, max_length=2000)
    footer_terms: str | None = Field(default=None, max_length=5000)
    standard_remark: str | None = Field(default=None, max_length=2000)


class TemplateOut(BaseModel):
    employee_code: str
    employee_name: str
    branch_code: str | None = None
    branch_name: str | None = None
    display_name: str | None = None
    phone: str | None = None
    logo_url: str | None = None
    bank_accounts: str | None = None
    footer_terms: str | None = None
    standard_remark: str | None = None
    updated_at: datetime | None = None
    month_end_preview: str          # วันท้ายบิลถ้าออกใบวันนี้ — ให้ DS เห็นว่า {month_end} จะกลายเป็นอะไร


def _tpl_out(emp: CurrentEmployee, t: QuotationTemplate | None) -> TemplateOut:
    return TemplateOut(
        employee_code=emp.employee_code, employee_name=emp.employee_name,
        branch_code=emp.branch_code, branch_name=emp.branch_name,
        display_name=t.display_name if t else None, phone=t.phone if t else None,
        logo_url=sales_extras_service.logo_url(t.logo_path) if t else None,
        bank_accounts=t.bank_accounts if t else None, footer_terms=t.footer_terms if t else None,
        standard_remark=t.standard_remark if t else None, updated_at=t.updated_at if t else None,
        month_end_preview=sales_extras_service.month_end(datetime.now().date()).strftime("%d/%m/%Y"),
    )


# ---------- ทะเบียนพนักงาน / บทบาท ----------
@router.get("/staff/employees", response_model=list[EmployeeOut])
def list_employees(db: Session = Depends(get_db), emp: CurrentEmployee = Depends(get_current_employee)):
    """รายชื่อสำหรับ dropdown "รหัส – ชื่อ" — วันที่ต่อ Employee API จะมาจากที่นั่นแทน"""
    return [EmployeeOut(user_id=r.user_id, employee_code=r.employee_code, employee_name=r.employee_name,
                        branch_code=r.branch_code, label=r.label)
            for r in get_employee_provider().list_employees(db)]


@router.get("/staff/roles", response_model=list[RoleOut])
def list_roles(emp: CurrentEmployee = Depends(get_current_employee)):
    return [RoleOut(code=k, name=v) for k, v in STAFF_ROLES.items()]


# ---------- ตะกร้า ----------
def _my_cart(db: Session, emp: CurrentEmployee, cart_id: str):
    user = db.get(User, emp.user_id)
    return sales_service.require_my_cart(db, user, cart_id)


@router.put("/sales/carts/{cart_id}/staff", response_model=CartOut)
def set_cart_staff(cart_id: str, body: StaffSetIn, db: Session = Depends(get_db),
                   emp: CurrentEmployee = Depends(get_current_employee)):
    cart = _my_cart(db, emp, cart_id)
    sales_extras_service.set_staff(db, emp, cart, body.role_code, body.user_id, body.employee_code)
    return cart_out(cart, db)


@router.put("/sales/carts/{cart_id}/remark", response_model=CartOut)
def set_cart_remark(cart_id: str, body: RemarkIn, db: Session = Depends(get_db),
                    emp: CurrentEmployee = Depends(get_current_employee)):
    cart = _my_cart(db, emp, cart_id)
    sales_extras_service.set_overall_remark(db, emp, cart, body.overall_remark)
    return cart_out(cart, db)


# ---------- template ----------
@router.get("/staff/quotation-template", response_model=TemplateOut)
def get_template(db: Session = Depends(get_db), emp: CurrentEmployee = Depends(get_current_employee)):
    """ของตัวเองเท่านั้น — ไม่มี endpoint ให้ระบุ id ของคนอื่น"""
    return _tpl_out(emp, sales_extras_service.get_template(db, emp))


@router.put("/staff/quotation-template", response_model=TemplateOut)
def save_template(body: TemplateIn, db: Session = Depends(get_db), emp: CurrentEmployee = Depends(get_current_employee)):
    return _tpl_out(emp, sales_extras_service.save_template(db, emp, body.model_dump(exclude_unset=True)))


@router.post("/staff/quotation-template/logo", response_model=TemplateOut)
async def upload_logo(file: UploadFile = File(...), db: Session = Depends(get_db),
                      emp: CurrentEmployee = Depends(get_current_employee)):
    return _tpl_out(emp, await sales_extras_service.save_logo(db, emp, file))


# ---------- export ----------
def _export_rows(q) -> tuple[dict, list[list], list[str]]:
    """ข้อมูลชุดเดียวกันทั้ง xlsx และ csv — หัวใบ · ตารางสินค้า · ท้ายใบ"""
    tp = q.template_snapshot or {}
    c = q.customer_snapshot or {}
    head = {
        "เลขที่ใบเสนอราคา": q.quotation_no,
        "วันที่ออก": q.issued_at.strftime("%d/%m/%Y %H:%M"),
        "ยืนราคาถึง": q.valid_until.strftime("%d/%m/%Y"),
        "วันที่ในเงื่อนไขท้ายบิล": month_end_of(q),
        "ลูกค้า": c.get("name") or "",
        "รหัสลูกค้า": c.get("sap_customer_no") or "",
        "พนักงานขาย": f"{tp.get('employee_code') or (q.sales.staff_code if q.sales else '') or ''} – "
                       f"{tp.get('employee_name') or (q.sales.name if q.sales else 'สั่งซื้อออนไลน์')}",
        "สาขา": " · ".join(x for x in [tp.get("branch_code") or (q.sales.branch_id if q.sales else None),
                                       tp.get("branch_name") or (q.sales.branch_name if q.sales else None)] if x),
    }
    for r in q.staff_snapshot or []:
        head[f"{r['role_code']} {r.get('role_name') or ''}".strip()] = f"{r['employee_code']} – {r['employee_name']}"
    lines = [["ลำดับ", "รหัสสินค้า", "รายการ", "จำนวน", "ราคาต่อหน่วย", "ส่วนลด", "จำนวนเงิน", "รับสินค้า", "หมายเหตุรายสินค้า"]]
    for i, l in enumerate(q.lines, 1):
        off = (Decimal(l.list_price) - Decimal(l.unit_price)) * l.qty
        lines.append([i, l.matnr, l.name + (f" ({l.variant})" if l.variant else ""), l.qty, float(l.unit_price),
                      float(off) if off > 0 else 0.0, float(l.line_total), SUPPLY_TXT.get(l.supply_mode, l.supply_mode or ""),
                      l.item_remark or ""])
    remark = "\n".join(x for x in [(q.overall_remark or (q.preso.note if q.preso else None) or "").strip(),
                                   (tp.get("standard_remark") or "").strip()] if x)
    foot = [
        ("รวมสินค้า", float(q.subtotal)),
        ("ส่วนลด", float(q.discount_total)),
        ("ค่าขนส่ง/ติดตั้ง", float(q.shipping_fee) + float(q.install_fee) - float(q.shipping_discount)),
        ("VAT 7% (รวมในยอด)", float(q.vat)),
        ("รวมสุทธิ", float(q.grand_total)),
    ]
    terms = _terms(q.valid_until.strftime("%d/%m/%Y"), month_end_of(q), tp.get("footer_terms"))
    return {"head": head, "foot": foot, "remark": remark, "bank": (tp.get("bank_accounts") or "").strip(),
            "logo": tp.get("logo_path")}, lines, terms


def _flat_rows(q) -> list[list]:
    """ใบทั้งใบเป็นแถวเรียงลงมา — ใช้ร่วมกันระหว่าง CSV กับ Google Sheets"""
    meta, lines, terms = _export_rows(q)
    rows: list[list] = [[k, v] for k, v in meta["head"].items()]
    rows.append([])
    rows.extend(lines)
    rows.append([])
    rows.extend(["", "", "", "", "", "", k, round(v, 2)] for k, v in meta["foot"])
    rows.append([])
    rows.append(["หมายเหตุ", meta["remark"]])
    if meta["bank"]:
        rows.append(["บัญชีสำหรับโอนชำระ", meta["bank"]])
    rows.append(["เงื่อนไข"])
    rows.extend([i, t] for i, t in enumerate(terms, 1))
    return rows


def _csv(q, bom: bool = True) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in _flat_rows(q):
        w.writerow([f"{v:.2f}" if isinstance(v, float) else v for v in r])
    # BOM ให้ Excel ภาษาไทยเปิดแล้วไม่เป็นตัวต่างดาว · IMPORTDATA ของ Google ไม่ต้องใช้ (จะติดมาเป็นอักษรล่องหน)
    return (("﻿" if bom else "") + buf.getvalue()).encode("utf-8")


def _xlsx(q) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    meta, lines, terms = _export_rows(q)
    wb = Workbook()
    ws = wb.active
    ws.title = q.quotation_no[:31]
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    row = 1
    logo = sales_extras_service.logo_file(meta["logo"])
    if logo:
        from openpyxl.drawing.image import Image as XLImage

        img = XLImage(str(logo))
        ratio = 60 / max(img.height, 1)        # สูงราว 60px ไม่บังข้อมูล
        img.height, img.width = int(img.height * ratio), int(img.width * ratio)
        ws.add_image(img, "A1")
        row = 5
    ws.cell(row=row, column=1, value="ใบเสนอราคา / Quotation").font = Font(bold=True, size=14)
    row += 1
    s_tp = q.template_snapshot or {}
    ws.cell(row=row, column=1, value=s_tp.get("company_name_th") or "")
    row += 2
    for k, v in meta["head"].items():
        ws.cell(row=row, column=1, value=k).font = bold
        ws.cell(row=row, column=2, value=v)
        row += 1
    row += 1
    fill = PatternFill("solid", fgColor="F2F2F0")
    for j, h in enumerate(lines[0], 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill = bold, fill
    row += 1
    for ln in lines[1:]:
        for j, v in enumerate(ln, 1):
            c = ws.cell(row=row, column=j, value=v)
            if j in (5, 6, 7):
                c.number_format = "#,##0.00"
            if j in (3, 9):
                c.alignment = wrap
        row += 1
    row += 1
    for k, v in meta["foot"]:
        ws.cell(row=row, column=6, value=k).font = bold
        c = ws.cell(row=row, column=7, value=v)
        c.number_format = "#,##0.00"
        if k == "รวมสุทธิ":
            c.font = bold
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="หมายเหตุ").font = bold
    ws.cell(row=row, column=2, value=meta["remark"]).alignment = wrap
    row += 1
    if meta["bank"]:
        ws.cell(row=row, column=1, value="บัญชีสำหรับโอนชำระ").font = bold
        ws.cell(row=row, column=2, value=meta["bank"]).alignment = wrap
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="เงื่อนไข").font = bold
    row += 1
    for i, t in enumerate(terms, 1):
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=t).alignment = wrap
        row += 1
    for col, wdt in zip("ABCDEFGHI", (22, 14, 44, 8, 14, 12, 16, 12, 30)):
        ws.column_dimensions[col].width = wdt
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@router.get("/quotations/{no}/export")
def export_quotation(no: str, format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"),
                     t: str | None = Query(default=None), bom: bool = Query(default=True),
                     db: Session = Depends(get_db), user: User | None = Depends(get_current_user_optional)):
    """ใช้สิทธิ์ชุดเดียวกับหน้าเอกสาร — ใครเปิดใบได้ export ได้"""
    q = quotation_service.get_quotation(db, no)
    quotation_service.check_access(q, user, t)
    if format == "csv":
        body, mime = _csv(q, bom), "text/csv; charset=utf-8"
    else:
        body, mime = _xlsx(q), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(body, media_type=mime,
                    headers={"Content-Disposition": f'attachment; filename="{q.quotation_no}.{format}"'})



class SheetOut(BaseModel):
    url: str


@router.post("/quotations/{no}/google-sheet", response_model=SheetOut)
def to_google_sheet(no: str, db: Session = Depends(get_db), emp: CurrentEmployee = Depends(get_current_employee)):
    """สร้าง Google Sheet ของใบนี้ใน Shared Drive ของบริษัท แล้วแชร์สิทธิ์แก้ไขให้คนที่กด

    ยังไม่ได้ตั้ง service account → 501 · หน้าเว็บจะใช้ทางสำรอง (สูตร IMPORTDATA) แทน
    """
    from app.integrations.google import sheets

    user = db.get(User, emp.user_id)
    q = quotation_service.get_quotation(db, no)
    quotation_service.check_access(q, user, None)
    try:
        url = sheets.create_sheet(f"{q.quotation_no} · {(q.customer_snapshot or {}).get('name') or ''}".strip(" ·"),
                                  _flat_rows(q), user.email if user else None)
    except sheets.SheetsNotConfigured:
        raise HTTPException(status_code=501, detail="ยังไม่ได้ตั้งค่า Google service account")
    except (sheets.SheetsError, httpx.HTTPError) as e:
        raise HTTPException(status_code=502, detail=str(e) or "ติดต่อ Google ไม่สำเร็จ")
    audit_service.log(db, user, "quotation.google_sheet", "quotation", q.quotation_no, {"url": url})
    db.commit()
    return SheetOut(url=url)
